from pptx import Presentation
from pptx.util import Inches, Pt
import os

def create_presentation():
    prs = Presentation()
    
    # Slide 1: Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "AI-Powered Personalized Blood Pressure Management Assistant"
    subtitle.text = "Implementation Updates & Final Evaluation\nMazvita Ziwira"

    # Slide 2: System Architecture Updates
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Enhanced Architecture & Security"
    tf = body_shape.text_frame
    tf.text = "User Authentication: Implemented robust login/registration on the frontend (React) with FastAPI backend."
    p = tf.add_paragraph()
    p.text = "Session Management: Integrated browser localStorage to persist user sessions."
    p = tf.add_paragraph()
    p.text = "Centralised API Routing: Unified api.js client automatically appends authenticated username for data isolation."
    p = tf.add_paragraph()
    p.text = "Persistent Storage: Backend uses a JSON-based persistence layer (storage.py) to permanently save readings."

    # Slide 3: Testing Results
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Rigorous Testing & Validation (Chapter 5)"
    tf = body_shape.text_frame
    tf.text = "Backend Validation: Pytest suite confirms 100% pass rate for data persistence and multi-user data isolation."
    p = tf.add_paragraph()
    p.text = "Frontend Validation: Vitest suite successfully validates the JNC 7 / ACC/AHA classification logic."
    
    img_path = r"C:\Users\ziwir\Smart AI Blood Pressure tracker\frontend\code_vitest_screenshot.png"
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(1), Inches(3.5), width=Inches(8))

    # Slide 4: Performance Evaluation
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "System Latency & Responsiveness (Chapter 6)"
    tf = body_shape.text_frame
    tf.text = "CRUD Operations: Read/Write operations are highly optimised, averaging under 100ms."
    p = tf.add_paragraph()
    p.text = "AI Engine: The /chat and /analysis endpoints demonstrate acceptable processing overhead for algorithmic classification."
    
    img_path2 = r"C:\Users\ziwir\Smart AI Blood Pressure tracker\frontend\chart_latencies_screenshot.png"
    if os.path.exists(img_path2):
        slide.shapes.add_picture(img_path2, Inches(1.5), Inches(3.5), height=Inches(3.5))

    # Slide 5: Trend Analysis
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Longitudinal Trend Analysis (Chapter 6)"
    tf = body_shape.text_frame
    tf.text = "The dashboard visualises multi-day blood pressure trends."
    p = tf.add_paragraph()
    p.text = "Validates Objective 7: Interactive trend charting to support user interpretation."
    
    img_path3 = r"C:\Users\ziwir\Smart AI Blood Pressure tracker\frontend\chart_trend_screenshot.png"
    if os.path.exists(img_path3):
        slide.shapes.add_picture(img_path3, Inches(1.5), Inches(3), height=Inches(4))

    # Slide 6: Personalised Recommendations
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Intelligent Health Coaching (Chapter 6)"
    tf = body_shape.text_frame
    tf.text = "The AI engine classifies users based on aggregated historical data."
    p = tf.add_paragraph()
    p.text = "Successfully correlates lifestyle tags (e.g., 'Stressed', 'High Salt') with BP spikes for personalised coaching."
    
    img_path4 = r"C:\Users\ziwir\Smart AI Blood Pressure tracker\frontend\recommendations_screenshot.png"
    if os.path.exists(img_path4):
        slide.shapes.add_picture(img_path4, Inches(1), Inches(3.5), width=Inches(8))

    prs.save("Updated_Final_Presentation.pptx")
    print("Successfully created Updated_Final_Presentation.pptx")

if __name__ == "__main__":
    create_presentation()
