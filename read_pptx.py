from pptx import Presentation

def extract_text(filename):
    prs = Presentation(filename)
    text = []
    for i, slide in enumerate(prs.slides):
        text.append(f"--- Slide {i+1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
        text.append("\n")
        
    with open("pptx_content.txt", "w", encoding="utf-8") as f:
        f.write("\n".join(text))
    print("Done writing to pptx_content.txt")

if __name__ == "__main__":
    extract_text(r"c:\Users\ziwir\Downloads\AI-Powered_Personalized_Blood_Pressure_Management_Assistant_Chapters_4-6_Presentation.pptx")
